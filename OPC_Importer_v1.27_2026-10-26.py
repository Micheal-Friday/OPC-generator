#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel -> editable VSDX, via an isolated invisible desktop Visio instance.

Windows + licensed/activated desktop Visio 2013 or newer required.
Install: python -m pip install openpyxl pywin32
Run: python OPC_Importer_v1.27_2026-10-26.py
Batch export and PDF merging follow case-insensitive alphabetical input filenames.
Excel need not be installed. No VBA macros or UserForms are used.
All inputs: empty Lane = 1; empty NextID = next nonempty Excel row, across lanes.
The final blank NextID remains unconnected. ID and all eight headers are required.
Control NoID is never inferred; explicit NextID values are preserved.
Input: .xlsx/.xlsm; formula cells must have saved calculated values.
"""
from __future__ import annotations

import os
import shutil
import unicodedata
import queue
import sys
import tempfile
import threading
import traceback
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

HEADERS = ['ID', 'Activity', 'InsideText', 'SideText', 'Lane',
           'LaneTitle', 'NextID', 'NoID']
PAGE_W, PAGE_H, MARGIN = 210 / 25.4, 297 / 25.4, 10 / 25.4
KINDS = {
    'operation': 'circle', 'عملیات': 'circle', 'circle': 'circle',
    'warehouse': 'triangle', 'انبار': 'triangle', 'storage': 'triangle',
    'triangle': 'triangle', 'qc': 'square', 'بازرسی': 'square',
    'inspection': 'square', 'square': 'square', 'control': 'diamond',
    'کنترل': 'diamond', 'decision': 'diamond', 'diamond': 'diamond',
}
DIGITS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize(value):
    return ''.join(text(value).lower().replace('ي', 'ی').replace('ك', 'ک')
                   .replace('\u200c', '').replace('\u200e', '')
                   .replace('\u200f', '').split())


def read_steps(ws):
    rows = ws.iter_rows(min_col=1, max_col=8)
    header = next(rows, ())
    if [text(c.value).lower() for c in header] != [h.lower() for h in HEADERS]:
        raise ValueError(f'{ws.title}: headers must be:\n' + ' | '.join(HEADERS))
    steps, by_id, titles, previous = [], {}, {}, {}
    for row_number, cells in enumerate(rows, 2):
        if any(c.data_type == 'e' for c in cells):
            raise ValueError(f'{ws.title}, row {row_number}: Excel error value.')
        values = [text(c.value) for c in cells]
        if not any(values):
            continue
        step_id, activity, inside, side, lane_text, title, nxt, no = values
        key = step_id.casefold()
        if not key or key in by_id:
            raise ValueError(f'{ws.title}, row {row_number}: missing/duplicate ID {step_id!r}.')
        kind = KINDS.get(normalize(activity))
        if not kind:
            raise ValueError(f'{ws.title}, row {row_number}: unknown Activity {activity!r}.')
        # Empty Lane means the first (or only) lane.
        lane_string = (lane_text or "1").translate(DIGITS)
        if not lane_string.isascii() or not lane_string.isdigit() or int(lane_string) < 1:
            raise ValueError(f'{ws.title}, row {row_number}: Lane must be a positive integer.')
        lane = int(lane_string)
        if titles.get(lane) and title and titles[lane] != title:
            raise ValueError(f'{ws.title}: inconsistent LaneTitle in lane {lane}.')
        titles[lane] = title or titles.get(lane, '')
        if no and kind != 'diamond':
            raise ValueError(f'{ws.title}: NoID is allowed only for Control: {step_id}.')
        s = dict(id=step_id, key=key, kind=kind, inside=inside, side=side,
                 lane=lane, next=nxt.casefold(), no=no.casefold(),
                 prior=previous.get(lane), level=1)
        previous[lane] = key
        by_id[key] = s
        steps.append(s)
    if not steps:
        raise ValueError(f'{ws.title}: no steps found.')
    if sorted(titles) != list(range(1, len(titles) + 1)):
        raise ValueError(f'{ws.title}: Lane numbers must be consecutive from 1.')
    # Blank NextID follows Excel row order, including across lanes.
    # Preserve explicit destinations; leave the final blank destination empty.
    for current, following in zip(steps, steps[1:]):
        if not current['next']:
            current['next'] = following['key']
    for s in steps:
        for field in ('next', 'no'):
            dest = s[field]
            if dest and dest not in by_id:
                raise ValueError(f'{ws.title}: unknown {field} ID {dest!r} for {s["id"]}.')
            if dest == s['key']:
                raise ValueError(f'{ws.title}: self-connection at {s["id"]}.')
    for _ in range(len(steps)):
        changed = False
        for s in steps:
            if s['prior']:
                required = by_id[s['prior']]['level'] + 1
                if s['level'] < required:
                    s['level'] = required
                    changed = True
            if s['next']:
                target = by_id[s['next']]
                if target['level'] <= s['level']:
                    target['level'] = s['level'] + 1
                    changed = True
        if not changed:
            break
    else:
        raise ValueError(f'{ws.title}: forward cycle or conflicting row order; use NoID for returns.')
    # Connections are validated above, but do not determine vertical placement.
    # Each lane starts at the same top position and follows its own Excel order.
    lane_rows = {}
    for step in steps:
        lane = step['lane']
        lane_rows[lane] = lane_rows.get(lane, 0) + 1
        step['level'] = lane_rows[lane]
    return steps, by_id, titles


def formulas(shape, **values):
    for name, value in values.items():
        shape.CellsU(name).FormulaU = str(value)


def apply_dubai(owner, doc, size=None):
    font_id = str(doc.Fonts.Item('Dubai').ID)
    for base in ('Char.Font', 'Char.ComplexScriptFont'):
        owner.CellsU(base).FormulaForceU = font_id
        index = 2
        while owner.CellExistsU(f'{base}{index}', 0):
            owner.CellsU(f'{base}{index}').FormulaForceU = font_id
            index += 1
    if size is not None:
        for base in ('Char.Size', 'Char.ComplexScriptSize'):
            owner.CellsU(base).FormulaForceU = f'{size:.6f} pt'
            index = 2
            while owner.CellExistsU(f'{base}{index}', 0):
                owner.CellsU(f'{base}{index}').FormulaForceU = f'{size:.6f} pt'
                index += 1


def font(shape, size, bold=False, direction='context'):
    apply_dubai(shape, shape.Document, size)
    for name, value in {'Char.Size': f'{size:.6f} pt', 'Char.Color': 'RGB(0,0,0)',
                        'Char.Style': '1' if bold else '0', 'Para.HorzAlign': '1',
                        'VerticalAlign': '1', 'LeftMargin': '1 pt',
                        'RightMargin': '1 pt', 'TopMargin': '1 pt',
                        'BottomMargin': '1 pt'}.items():
        shape.CellsU(name).FormulaU = value
    apply_text_direction(shape, direction)


def detect_text_direction(value, fallback='rtl'):
    # Ignore numbers, punctuation, whitespace and bidi control characters.
    # Use the first actual letter with a strong Unicode bidi direction.
    for char in str(value):
        if not unicodedata.category(char).startswith('L'):
            continue
        bidi = unicodedata.bidirectional(char)
        if bidi in ('R', 'AL'):
            return 'rtl'
        if bidi == 'L':
            return 'ltr'
    return fallback


def set_paragraph_direction(owner, direction='rtl'):
    # Both Style and Shape objects support named paragraph cell checks.
    if direction not in ('rtl', 'ltr'):
        raise ValueError('Resolved text direction must be rtl or ltr.')
    flag = '1' if direction == 'rtl' else '0'
    owner.CellsU('Para.Flags').FormulaForceU = flag
    index = 2
    while owner.CellExistsU(f'Para.Flags{index}', 0):
        owner.CellsU(f'Para.Flags{index}').FormulaForceU = flag
        index += 1


def apply_text_direction(shape, direction='context'):
    if direction != 'context':
        set_paragraph_direction(shape, direction)
        return
    value = str(shape.Text)
    resolved = detect_text_direction(value)
    set_paragraph_direction(shape, resolved)
    # Paragraphs can share a ShapeSheet formatting row. Apply differing
    # directions to character ranges so Visio creates the necessary rows.
    offset = 0
    for paragraph in value.splitlines(keepends=True):
        length = len(paragraph.encode('utf-16-le')) // 2
        body = paragraph.rstrip('\r\n\u2028\u2029')
        paragraph_direction = detect_text_direction(body, resolved)
        if body and paragraph_direction != resolved:
            import pythoncom
            chars = shape.Characters
            chars.Begin = offset
            chars.End = offset + len(body.encode('utf-16-le')) // 2
            # ParaProps is an indexed COM property; visFlags = 13.
            dispid = chars._oleobj_.GetIDsOfNames('ParaProps')
            chars._oleobj_.Invoke(dispid, 0, pythoncom.DISPATCH_PROPERTYPUT,
                                 0, 13, 1 if paragraph_direction == 'rtl' else 0)
        offset += length


def set_document_direction(doc, direction='context'):
    # Start with LTR formatting; each generated text gets its own direction.
    if direction not in ('context', 'rtl', 'ltr'):
        raise ValueError('Text direction must be context, rtl or ltr.')
    style = doc.Styles.Item(doc.DefaultTextStyle)
    try:
        apply_dubai(style, doc)
    except Exception as exc:
        raise RuntimeError('فونت Dubai باید روی ویندوز نصب باشد.') from exc
    set_paragraph_direction(style, 'ltr' if direction == 'context' else direction)
    if style.CellExistsU('Char.RTLText', 0):
        style.CellsU('Char.RTLText').FormulaForceU = '0'
    doc.DefaultTextStyle = style.Name


def set_ltr_character_runs(shape):
    # Clear inherited RTL character formatting independently of paragraph flags.
    index = 1
    while True:
        cell = 'Char.RTLText' + (str(index) if index > 1 else '')
        if not shape.CellExistsU(cell, 0):
            break
        shape.CellsU(cell).FormulaForceU = '0'
        index += 1


def set_english_text_language(shape):
    # Windows LCID 1033 = en-US. Apply locally to every character run.
    index = 1
    while True:
        cell = 'Char.LangID' + (str(index) if index > 1 else '')
        if index > 1 and not shape.CellExistsU(cell, 0):
            break
        shape.CellsU(cell).FormulaForceU = '1033'
        index += 1


def textbox(pg, left, bottom, right, top, value, size, bold=False, direction='context'):
    s = pg.DrawRectangle(left, bottom, right, top)
    # Set direction BEFORE assigning text so the document's RTL default does
    # not seed RTL formatting into a Latin filename containing digits/hyphens.
    resolved = detect_text_direction(value) if direction == 'context' else direction
    set_paragraph_direction(s, resolved)
    if direction == 'ltr':
        set_ltr_character_runs(s)
    if direction == 'ltr' and str(value).isascii():
        set_english_text_language(s)
    s.Text = value
    if direction == 'ltr' and str(value).isascii():
        set_english_text_language(s)
    formulas(s, LinePattern=0, FillPattern=0)
    font(s, size, bold, direction)
    if direction == 'ltr':
        set_ltr_character_runs(s)
    return s


def symbol(pg, kind, x, y, half):
    if kind == 'circle':
        s = pg.DrawOval(x-half, y-half, x+half, y+half)
    elif kind == 'square':
        s = pg.DrawRectangle(x-half, y-half, x+half, y+half)
    else:
        # Build polygon geometry with scalar ShapeSheet formulas.
        # Avoid SAFEARRAY marshalling differences in DrawPolyline wrappers.
        s = pg.DrawRectangle(x-half, y-half, x+half, y+half)
        section = 10  # visSectionFirstComponent
        s.DeleteSection(section)
        s.AddSection(section)
        s.AddRow(section, 0, 137)  # visTagComponent
        if kind == 'triangle':
            points = [(0, 1), (1, 1), (.5, 0), (0, 1)]
        else:
            points = [(.5, 1), (1, .5), (.5, 0), (0, .5), (.5, 1)]
        for i, (px, py) in enumerate(points):
            row = s.AddRow(section, i + 1, 138 if i == 0 else 139)
            s.CellsSRC(section, row, 0).FormulaU = f"Width*{px:g}"
            s.CellsSRC(section, row, 1).FormulaU = f"Height*{py:g}"
        s.CellsU('Geometry1.NoFill').FormulaU = '0'
        s.CellsU('Geometry1.NoLine').FormulaU = '0'
    formulas(s, FillPattern=1, FillForegnd='RGB(255,255,255)',
             LinePattern=1, LineColor='RGB(0,0,0)', LineWeight='0.75 pt')
    return s


def document_display_name(path):
    # Remove ONLY Unicode directional controls, never visible filename text.
    # These can be present in names copied from RTL applications.
    controls = {'\u061c', '\u200e', '\u200f', '\u202a', '\u202b',
                '\u202c', '\u202d', '\u202e', '\u2066', '\u2067',
                '\u2068', '\u2069'}
    return ''.join(char for char in Path(path).stem if char not in controls)


def legend(pg, document_name, process_name, direction='context'):
    """Full-width footer: process | document | two-row symbol guide."""
    left, right = MARGIN, PAGE_W - MARGIN
    bottom, top = MARGIN, MARGIN + 1.05
    guide_left = right - 2.9
    middle = left + (guide_left - left) / 2
    frame = pg.DrawRectangle(left, bottom, right, top)
    formulas(frame, FillPattern=0, LinePattern=1, LineColor='RGB(0,0,0)', LineWeight='0.75 pt')
    for x in (middle, guide_left):
        divider = pg.DrawLine(x, bottom, x, top)
        formulas(divider, LinePattern=1, LineColor='RGB(0,0,0)',
                 LineWeight='0.75 pt', BeginArrow=0, EndArrow=0)
    # Editable text uses the selected paragraph direction.
    for x1, x2, title, value in (
        (left, middle, 'نام فرایند', process_name),
        (middle, guide_left, 'نام سند', document_name),
    ):
        textbox(pg, x1+.08, top-.29, x2-.08, top-.05, title, 9, True, direction)
        name_shape = textbox(pg, x1+.08, bottom+.10, x2-.08, top-.34,
            value, 9,
            direction='ltr' if title == 'نام سند' else direction)
        if title == 'نام سند':
            set_paragraph_direction(name_shape, 'ltr')
    textbox(pg, guide_left+.08, top-.27, right-.08, top-.03, 'راهنما', 9, True, direction)
    # Right-to-left reading order: square/triangle, then circle/diamond.
    items = [('square', 'بازرسی'), ('triangle', 'انبار'),
             ('circle', 'عملیات'), ('diamond', 'کنترل')]
    cell_width = (right-guide_left-.16)/2
    for i, (kind, label) in enumerate(items):
        row, column = divmod(i, 2)
        cell_left = guide_left+.08+(1-column)*cell_width
        y = top-.46-row*.34
        x = cell_left+.14
        symbol(pg, kind, x, y, .12)
        textbox(pg, x+.12+.05, y-.14, cell_left+cell_width-.04,
                y+.14, label, 9, True, direction)


def draw_page(app, pg, model, document_name, process_name, direction='context'):
    steps, by_id, titles = model
    pg.AutoSize = False
    formulas(pg.PageSheet, PageScale='1 mm', DrawingScale='1 mm',
             PageWidth='210 mm', PageHeight='297 mm', PaperKind=9,
             PrintPageOrientation=1, OnPage='TRUE', PagesX=1, PagesY=1, RouteStyle=1)
    width = len(titles)*2.35 + (.45 if any(s['no'] for s in steps) else 0)
    height = .45 + .18 + .6 + (max(s['level'] for s in steps)-1)*.9
    fit = min(1, (PAGE_W-2*MARGIN)/width, (PAGE_H-2*MARGIN-1.20)/height)
    half = .3*fit
    top = PAGE_H-MARGIN
    # No worksheet title at the top, as requested.
    for lane, title in sorted(titles.items()):
        left = MARGIN+(lane-1)*2.35*fit
        textbox(pg, left, top-.45*fit, left+2.17*fit, top, title, 11*fit, True, direction)
    for s in steps:
        x = MARGIN+(s['lane']-1)*2.35*fit+half
        y = top-(.45+.18)*fit-half-(s['level']-1)*.9*fit
        s['x'] = x
        shape = symbol(pg, s['kind'], x, y, half)
        shape.Text = s['inside']
        font(shape, 11*fit, direction=direction)
        s['shape'] = shape
        textbox(pg, x+half+.12*fit, y-half, x+half+1.57*fit, y+half, s['side'], 10*fit, direction=direction)
    legend(pg, document_name, process_name, direction)
    for s in steps:
        for field in ('next', 'no'):
            if not s[field]:
                continue
            target = by_id[s[field]]
            link = pg.Drop(app.ConnectorToolDataObject, 0, 0)
            formulas(link, ShapeRouteStyle=1, LinePattern=1, LineColor='RGB(0,0,0)',
                     LineWeight='0.75 pt', BeginArrow=0, FillPattern=0)
            is_no = field == 'no'
            arrow = is_no or s['kind'] == 'diamond' or s['lane'] != target['lane']
            formulas(link, EndArrow=4 if arrow else 0)
            link.CellsU('BeginX').GlueToPos(s['shape'], 1 if is_no else .5, .5 if is_no else 0)
            if is_no or s['lane'] == target['lane']:
                link.CellsU('EndX').GlueToPos(target['shape'], .5, 1)
            else:
                link.CellsU('EndX').GlueToPos(target['shape'], 0 if s['x'] < target['x'] else 1, .5)
            if s['kind'] == 'diamond':
                link.Text = 'خیر' if is_no else 'بله'
                font(link, 9*fit, direction=direction)


def export_visio(excel_path, selected_names, output, report, direction='context', export_pdf=False, save_visio=True):
    from openpyxl import load_workbook
    import pythoncom
    from win32com.client import DispatchEx

    # Validate all selected sheets before launching Visio.
    report('در حال بررسی اطلاعات اکسل…')
    book = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        models = [(name, read_steps(book[name])) for name in selected_names]
    finally:
        book.close()

    pythoncom.CoInitialize()
    app = doc = None
    temp_path = None
    temp_pdf = None
    try:
        report('در حال اجرای Visio در پس‌زمینه…')
        try:
            app = DispatchEx('Visio.InvisibleApp')
        except Exception as exc:
            raise RuntimeError('Visio دسکتاپ روی ویندوز باید نصب و فعال باشد.\n' + str(exc)) from exc
        app.Visible = False
        doc = app.Documents.Add('')
        set_document_direction(doc, direction)
        for number, (name, model) in enumerate(models, 1):
            report(f'ساخت صفحه {number} از {len(models)} — {name}')
            pg = doc.Pages.Item(1) if number == 1 else doc.Pages.Add()
            pg.Name = str(number)
            draw_page(app, pg, model, document_display_name(excel_path), name, direction)
        report('در حال ذخیره فایل Visio…')
        output = Path(output).resolve()
        if save_visio:
            fd, temp_name = tempfile.mkstemp(prefix='OPC_', suffix='.vsdx', dir=output.parent)
            os.close(fd)
            temp_path = Path(temp_name)
            temp_path.unlink()  # SaveAs writes a new file without an overwrite prompt.
            doc.SaveAs(str(temp_path))
        if export_pdf:
            report('در حال ساخت PDF از همه صفحات…')
            fd, pdf_name = tempfile.mkstemp(prefix='OPC_', suffix='.pdf', dir=output.parent)
            os.close(fd)
            temp_pdf = Path(pdf_name)
            temp_pdf.unlink()
            # PDF=1, print quality=1, all foreground pages=0.
            doc.ExportAsFixedFormat(1, str(temp_pdf), 1, 0)
            if not temp_pdf.exists() or temp_pdf.stat().st_size == 0:
                raise RuntimeError('Visio did not create the PDF file.')
        doc.Saved = True
        doc.Close()
        doc = None
        if save_visio:
            if not temp_path.exists() or temp_path.stat().st_size == 0:
                raise RuntimeError('Visio did not create an output file.')
            os.replace(temp_path, output)
            temp_path = None
        if temp_pdf is not None:
            try:
                os.replace(temp_pdf, output.with_suffix('.pdf'))
            except OSError as exc:
                raise RuntimeError(f'Visio saved: {output}\nPDF could not be saved: {exc}') from exc
            temp_pdf = None
    finally:
        if doc is not None:
            try:
                doc.Saved = True
                doc.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass
        if temp_pdf is not None:
            try:
                temp_pdf.unlink(missing_ok=True)
            except OSError:
                pass
        pythoncom.CoUninitialize()


def filename_sort_key(path):
    """Sort by filename rather than folder, ignoring letter case."""
    return unicodedata.normalize('NFC', Path(path).name).casefold()


def plan_batch_outputs(jobs, folder, export_pdf=False):
    planned, used = [], set()
    for path, names in sorted(jobs, key=lambda job: filename_sort_key(job[0])):
        output = Path(folder) / (Path(path).stem + '.vsdx')
        key = str(output.resolve()).casefold()
        if key in used:
            raise OSError('دو فایل ورودی نام یکسان دارند؛ برای هرکدام جداگانه خروجی بگیرید: ' + output.name)
        used.add(key)
        planned.append((path, names, str(output)))
    return planned


def inspect_workbook(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        results = {}
        for ws in book.worksheets:
            try:
                read_steps(ws)
                results[ws.title] = ''
            except Exception as exc:
                results[ws.title] = str(exc)
        return results
    finally:
        book.close()


def run_batch(jobs, report, direction='context', exporter=export_visio, export_pdf=False):
    jobs = sorted(jobs, key=lambda job: filename_sort_key(job[0]))
    successes, failures = [], []
    for number, (path, names, output) in enumerate(jobs, 1):
        prefix = f'[{number}/{len(jobs)}] {Path(path).name} — '
        try:
            exporter(path, names, output, lambda msg, p=prefix: report(p + msg), direction=direction, export_pdf=export_pdf)
            successes.append(str(output) + ('\n' + str(Path(output).with_suffix('.pdf')) if export_pdf else ''))
        except Exception as exc:
            if sys.stderr is not None:
                traceback.print_exc()
            failures.append((path, str(exc)))
    return successes, failures


def merge_pdf_files(paths, output):
    """Merge ordered PDFs without changing originals; publish only a complete file."""
    from pypdf import PdfWriter
    output = Path(output).resolve()
    paths = [Path(p).resolve() for p in paths]
    if not paths:
        raise ValueError('No PDFs to merge.')
    if any(str(p).casefold() == str(output).casefold() for p in paths):
        raise ValueError('Merged PDF must have a separate filename.')
    fd, name = tempfile.mkstemp(prefix='OPC_merge_', suffix='.pdf', dir=output.parent)
    os.close(fd)
    temp = Path(name)
    writer = PdfWriter()
    try:
        for path in paths:
            writer.append(str(path))
        writer.write(str(temp))
        writer.close()
        os.replace(temp, output)
    finally:
        writer.close()
        temp.unlink(missing_ok=True)


def publish_file(source, destination):
    destination = Path(destination)
    fd, name = tempfile.mkstemp(prefix='OPC_publish_', suffix=destination.suffix, dir=destination.parent)
    os.close(fd)
    try:
        shutil.copyfile(source, name)
        os.replace(name, destination)
    finally:
        Path(name).unlink(missing_ok=True)


def process_outputs(jobs, report, direction, save_visio, individual_pdf, merge_output, skipped=None, exporter=export_visio):
    # Sort source names before staging: numbered temporary PDFs must retain
    # this order when merged (lexical sorting would put 10.pdf before 2.pdf).
    jobs = sorted(jobs, key=lambda job: filename_sort_key(job[0]))
    skipped = skipped or set()
    successes, failures, pdfs = [], [], []
    need_pdf = individual_pdf or bool(merge_output)
    with tempfile.TemporaryDirectory(prefix='OPC_batch_') as scratch:
        for i, (path, names, output) in enumerate(jobs):
            staged = Path(scratch) / f'{i}.vsdx'
            saved = []
            try:
                report(f'[{i+1}/{len(jobs)}] {Path(path).name}')
                exporter(path, names, staged, report, direction=direction,
                             export_pdf=need_pdf, save_visio=save_visio)
                if save_visio and str(Path(output).resolve()).casefold() not in skipped:
                    publish_file(staged, output)
                    saved.append(str(output))
                if individual_pdf and str(Path(output).with_suffix('.pdf').resolve()).casefold() not in skipped:
                    target = Path(output).with_suffix('.pdf')
                    publish_file(staged.with_suffix('.pdf'), target)
                    saved.append(str(target))
                if need_pdf:
                    pdfs.append(staged.with_suffix('.pdf'))
                successes.append('\n'.join(saved) if saved else Path(path).name + (' — آماده برای ادغام' if merge_output else ' — ذخیره خروجی لغو شد'))
            except Exception as exc:
                failures.append((path, str(exc) + ('\nذخیره‌شده: ' + ', '.join(saved) if saved else '')))
        note = ''
        if merge_output:
            if failures:
                note = 'ادغام انجام نشد؛ ساخت بعضی فایل‌ها ناموفق بود.'
            else:
                try:
                    report('در حال ادغام PDFها…')
                    merge_pdf_files(pdfs, merge_output)
                    note = 'PDF تجمیعی ذخیره شد:\n' + str(merge_output)
                except Exception as exc:
                    note = 'خطا در ذخیره PDF تجمیعی:\n' + str(exc)
    return successes, failures, note


def export_existing_visio(path, names, output, report, **options):
    """Read an existing drawing as-is; do not alter its font, layout or pages."""
    import pythoncom
    from win32com.client import DispatchEx
    pythoncom.CoInitialize()
    app = doc = None
    try:
        app = DispatchEx('Visio.InvisibleApp')
        app.Visible = False
        # Read-only, hidden, macros disabled, no workspace restoration.
        doc = app.Documents.OpenEx(str(Path(path).resolve()), 2 | 64 | 128 | 256)
        target = Path(output).with_suffix('.pdf')
        report('تبدیل به PDF: ' + Path(path).name)
        doc.ExportAsFixedFormat(1, str(target), 1, 0)
        if not target.exists() or target.stat().st_size == 0:
            raise RuntimeError('Visio did not create the PDF file.')
    finally:
        if doc is not None:
            try:
                doc.Saved = True
                doc.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


class ImportWindow(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('OPC — Excel to Visio | v1.27')
        self.geometry('740x620')
        self.minsize(560, 440)
        self.path = tk.StringVar(value='هنوز فایلی انتخاب نشده است')
        self.status = tk.StringVar(value='فایل‌های اکسل را اضافه کنید؛ شیت‌های هر فایل را جداگانه تیک بزنید.')
        self.books = {}
        self.validation = {}
        self.file_errors = {}
        self.summary = tk.StringVar(value="کل: ۰ | آماده: ۰ | دارای خطا: ۰ | بدون انتخاب: ۰")
        self.checks = []
        self.group_buttons = []
        self.busy = False
        self.save_visio = tk.BooleanVar(value=True)
        self.export_pdf = tk.BooleanVar(value=False)
        self.merge_pdf = tk.BooleanVar(value=False)
        self.events = queue.Queue()
        bar = ttk.Frame(self, padding=12)
        bar.pack(fill='x')
        ttk.Entry(bar, textvariable=self.path, state='readonly').pack(side='left', fill='x', expand=True)
        self.browse = ttk.Button(bar, text='افزودن فایل‌های اکسل', command=self.choose_file)
        self.browse.pack(side='right', padx=(8, 0))
        ttk.Label(self, text='هر فایل اکسل → یک فایل Visio مستقل', anchor='e').pack(fill='x', padx=12)
        ttk.Label(self, textvariable=self.summary, anchor='e').pack(fill='x', padx=12, pady=6)
        self.recheck_button = ttk.Button(self, text='بررسی مجدد فایل‌ها', command=self.recheck)
        self.recheck_button.pack(anchor='e', padx=12)
        self.exporter_button = ttk.Button(self, text='تبدیل فایل Visio به PDF', command=self.open_exporter)
        self.exporter_button.pack(anchor='e', padx=12, pady=4)
        holder = ttk.Frame(self, padding=12)
        holder.pack(fill='both', expand=True)
        self.canvas = tk.Canvas(holder, highlightthickness=0)
        scroll = ttk.Scrollbar(holder, orient='vertical', command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=scroll.set)
        scroll.pack(side='right', fill='y')
        self.canvas.pack(side='left', fill='both', expand=True)
        self.list_frame = ttk.Frame(self.canvas)
        window_id = self.canvas.create_window((0, 0), window=self.list_frame, anchor='nw')
        self.list_frame.bind('<Configure>', lambda e: self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.canvas.bind('<Configure>', lambda e: self.canvas.itemconfigure(window_id, width=e.width))
        self.bind_all('<MouseWheel>', self.scroll_file_list, add='+')
        self.bind_all('<Button-4>', self.scroll_file_list, add='+')
        self.bind_all('<Button-5>', self.scroll_file_list, add='+')
        self.visio_check = ttk.Checkbutton(self, text='\u202bذخیره فایل \u202aVisio\u202c\u202c', variable=self.save_visio)
        self.visio_check.pack(anchor='e', padx=12, pady=(4, 0))
        self.pdf_check = ttk.Checkbutton(self, text='\u202bذخیره \u202aPDF\u202c تکی\u202c', variable=self.export_pdf)
        self.pdf_check.pack(anchor='e', padx=12, pady=(4, 0))
        self.merge_check = ttk.Checkbutton(self, text='\u202bادغام \u202aPDF\u202cها در یک فایل جداگانه\u202c',
            variable=self.merge_pdf)
        self.merge_check.pack(anchor='e', padx=12, pady=(4, 0))
        buttons = ttk.Frame(self, padding=12)
        buttons.pack(fill='x')
        self.all_button = ttk.Button(buttons, text='انتخاب همه شیت‌ها', command=lambda: self.select_all(True))
        self.all_button.pack(side='left')
        self.none_button = ttk.Button(buttons, text='حذف تیک‌ها', command=lambda: self.select_all(False))
        self.none_button.pack(side='left', padx=6)
        self.create_button = ttk.Button(buttons, text='ساخت خروجی‌ها', command=self.create)
        self.create_button.pack(side='right')
        ttk.Label(self, textvariable=self.status, wraplength=690, anchor='e').pack(fill='x', padx=12, pady=12)
        self.protocol('WM_DELETE_WINDOW', self.close)
        self.after(150, self.poll)

    def open_exporter(self):
        dialog = tk.Toplevel(self)
        dialog.title('Visio → PDF')
        dialog.transient(self)
        dialog.grab_set()
        paths = []
        label = tk.StringVar(value='فایل Visio انتخاب نشده است')
        single = tk.BooleanVar(value=True)
        merged = tk.BooleanVar(value=False)
        def browse():
            chosen = filedialog.askopenfilenames(parent=dialog, title='انتخاب فایل‌های Visio',
                filetypes=[('Visio drawings', '*.vsdx *.vsdm *.vsd')])
            if chosen:
                paths[:] = sorted(dict.fromkeys(chosen), key=filename_sort_key)
                label.set('\n'.join(Path(p).name for p in paths))
        ttk.Button(dialog, text='انتخاب فایل‌های Visio', command=browse).pack(padx=16, pady=12)
        ttk.Label(dialog, textvariable=label, wraplength=560).pack(padx=16)
        ttk.Checkbutton(dialog, text='PDF تکی', variable=single).pack(anchor='e', padx=16, pady=8)
        ttk.Checkbutton(dialog, text='PDF تجمیعی', variable=merged).pack(anchor='e', padx=16)
        def start():
            if not paths or not (single.get() or merged.get()):
                messagebox.showinfo('انتخاب خروجی', 'فایل و حداقل یک نوع PDF را انتخاب کنید.', parent=dialog)
                return
            if merged.get():
                try:
                    import pypdf
                except ImportError:
                    messagebox.showerror('کتابخانه لازم', 'python -m pip install pypdf', parent=dialog)
                    return
            folder = filedialog.askdirectory(parent=dialog, title='پوشه خروجی PDF')
            if not folder:
                return
            try:
                jobs = plan_batch_outputs([(p, []) for p in paths], folder)
            except OSError as exc:
                messagebox.showerror('نام تکراری', str(exc), parent=dialog)
                return
            merged_path = None
            if merged.get():
                merged_path = filedialog.asksaveasfilename(parent=dialog, title='ذخیره PDF تجمیعی',
                    initialdir=folder, initialfile='Merged.pdf', defaultextension='.pdf',
                    filetypes=[('PDF', '*.pdf')], confirmoverwrite=False)
                if not merged_path:
                    return
                if Path(merged_path).suffix.lower() != '.pdf':
                    messagebox.showerror('پسوند', 'پسوند باید .pdf باشد.', parent=dialog)
                    return
                if single.get() and any(str(Path(j[2]).with_suffix('.pdf').resolve()).casefold() == str(Path(merged_path).resolve()).casefold() for j in jobs):
                    messagebox.showerror('نام فایل', 'نام PDF تجمیعی باید با فایل‌های تکی متفاوت باشد.', parent=dialog)
                    return
            targets = [Path(j[2]).with_suffix('.pdf') for j in jobs] if single.get() else []
            if merged_path:
                targets.append(Path(merged_path))
            skipped = set()
            for target in targets:
                if target.exists() and not self.confirm_replace(target):
                    skipped.add(str(target.resolve()).casefold())
                    dialog.grab_set()
            if merged_path and str(Path(merged_path).resolve()).casefold() in skipped:
                merged_path = None
            if all(str(p.resolve()).casefold() in skipped for p in targets):
                return
            individual = single.get()
            dialog.destroy()
            self.busy = True
            self.set_enabled(False)
            def worker():
                try:
                    result = process_outputs(jobs, lambda msg: self.events.put(('status', msg)),
                        'context', False, individual, merged_path, skipped, exporter=export_existing_visio)
                    self.events.put(('done', result))
                except Exception as exc:
                    self.events.put(('error', str(exc)))
            threading.Thread(target=worker, daemon=False).start()
        ttk.Button(dialog, text='تبدیل به PDF', command=start).pack(padx=16, pady=16)

    def scroll_file_list(self, event):
        # Wheel events can originate from labels/checkbuttons inside the canvas.
        # Route only when the pointer is inside this list, not other dialogs.
        widget = self.winfo_containing(event.x_root, event.y_root)
        while widget is not None and widget is not self.canvas:
            widget = getattr(widget, 'master', None)
        if widget is None:
            return
        bounds = self.canvas.bbox('all')
        if not bounds or bounds[3] - bounds[1] <= self.canvas.winfo_height():
            return 'break'
        number = getattr(event, 'num', None)
        delta = getattr(event, 'delta', 0)
        if number == 4:
            units = -1
        elif number == 5:
            units = 1
        elif delta:
            units = (-1 if delta > 0 else 1) * max(1, int(abs(delta) / 120))
        else:
            return
        self.canvas.yview_scroll(units, 'units')
        return 'break'

    def choose_file(self):
        paths = filedialog.askopenfilenames(parent=self, title='انتخاب یک یا چند فایل اکسل',
                                          filetypes=[('Excel', '*.xlsx *.xlsm')])
        if not paths:
            return
        from openpyxl import load_workbook
        errors = []
        for raw_path in paths:
            path = str(Path(raw_path).resolve())
            if any(os.path.normcase(path) == os.path.normcase(p) for p in self.books):
                continue
            try:
                checks = inspect_workbook(path)
                self.validation[path] = checks
                self.books[path] = [(name, tk.BooleanVar(value=False)) for name in checks]
            except Exception as exc:
                self.books[path] = []
                self.file_errors[path] = str(exc)
                errors.append(Path(path).name + ': ' + str(exc))
        self.render_books()
        if errors:
            messagebox.showerror('خطا در خواندن فایل', '\n\n'.join(errors), parent=self)
        self.status.set('از هر فایل حداقل یک شیت انتخاب کنید؛ فایل بدون تیک تبدیل نمی‌شود.')

    def render_books(self):
        for widget in self.list_frame.winfo_children():
            widget.destroy()
        self.checks.clear()
        self.group_buttons.clear()
        self.path.set(f'{len(self.books)} فایل اکسل در فهرست')
        self.update_summary()
        for path, choices in sorted(self.books.items(), key=lambda item: filename_sort_key(item[0])):
            group = ttk.LabelFrame(self.list_frame, text=Path(path).name, padding=8)
            group.pack(fill='x', padx=4, pady=6)
            ttk.Label(group, text=str(Path(path).parent), wraplength=610).pack(fill='x')
            if path in self.file_errors:
                ttk.Label(group, text=self.file_errors[path], foreground='red', wraplength=610).pack(fill='x')
            controls = ttk.Frame(group)
            controls.pack(fill='x', pady=4)
            for title, command in [
                ('انتخاب شیت‌های این فایل', lambda p=path: self.select_all(True, p)),
                ('حذف تیک‌ها', lambda p=path: self.select_all(False, p)),
                ('حذف فایل', lambda p=path: self.remove_file(p)),
            ]:
                button = ttk.Button(controls, text=title, command=command)
                button.pack(side='left', padx=2)
                self.group_buttons.append(button)
            for name, variable in choices:
                widget = ttk.Checkbutton(group, text=name + (' — خطا' if self.validation.get(path, {}).get(name) else ' — معتبر'), variable=variable, command=self.update_summary)
                widget.pack(fill='x', padx=6, pady=3)
                self.checks.append(widget)
                error = self.validation.get(path, {}).get(name)
                if error:
                    ttk.Label(group, text=error, foreground="red", wraplength=610).pack(fill="x", padx=20)

    def remove_file(self, path):
        del self.books[path]
        self.validation.pop(path, None)
        self.file_errors.pop(path, None)
        self.render_books()

    def select_all(self, value, path=None):
        groups = [self.books[path]] if path is not None else self.books.values()
        for choices in groups:
            for _, variable in choices:
                variable.set(value)
        self.update_summary()

    def update_summary(self):
        ready = bad = unselected = 0
        for path, choices in self.books.items():
            names = [name for name, var in choices if var.get()]
            checks = self.validation.get(path, {})
            # Inventory errors are independent of sheet selection.
            if path in self.file_errors or any(checks.values()):
                bad += 1
            if not names:
                unselected += 1
            elif path not in self.file_errors and all(name in checks and not checks[name] for name in names):
                ready += 1
        self.summary.set(f'کل: {len(self.books)} | آماده ساخت: {ready} | دارای خطا (همه فایل‌ها): {bad} | بدون انتخاب: {unselected}')

    def recheck(self):
        for path in self.books:
            selected = {n for n, v in self.books[path] if v.get()}
            try:
                checks = inspect_workbook(path)
                self.validation[path] = checks
                self.books[path] = [(n, tk.BooleanVar(value=n in selected)) for n in checks]
                self.file_errors.pop(path, None)
            except Exception as exc:
                self.file_errors[path] = str(exc)
        self.render_books()

    def confirm_replace(self, path):
        dialog = tk.Toplevel(self)
        dialog.title('فایل موجود است')
        dialog.transient(self)
        dialog.grab_set()
        accepted = [False]
        ttk.Label(dialog, text='این فایل موجود است؛ جایگزین شود؟\n' + str(path),
                  wraplength=560, padding=16).pack()
        buttons = ttk.Frame(dialog, padding=12)
        buttons.pack()
        def replace():
            accepted[0] = True
            dialog.destroy()
        ttk.Button(buttons, text='Replace', command=replace).pack(side='left', padx=8)
        ttk.Button(buttons, text='Cancel', command=dialog.destroy).pack(side='left', padx=8)
        self.wait_window(dialog)
        return accepted[0]

    def create(self):
        self.recheck()
        save_visio = self.save_visio.get()
        individual_pdf = self.export_pdf.get()
        export_pdf = individual_pdf or self.merge_pdf.get()
        if not (save_visio or export_pdf):
            messagebox.showinfo('انتخاب خروجی', 'حداقل یک نوع خروجی را انتخاب کنید.', parent=self)
            return
        if self.merge_pdf.get():
            try:
                import pypdf
            except ImportError:
                messagebox.showerror('نیاز به نصب کتابخانه',
                    'برای ادغام PDF این دستور را اجرا کنید:\npython -m pip install pypdf', parent=self)
                return
        jobs = []
        for path, choices in self.books.items():
            names = [name for name, variable in choices if variable.get()]
            if names:
                jobs.append((path, names))
        invalid = [path for path, names in jobs if path in self.file_errors or any(self.validation.get(path, {}).get(n) for n in names)]
        if invalid:
            messagebox.showerror('خطای ورودی', 'خطاهای شیت‌های انتخاب‌شده را اصلاح کنید یا تیک آن‌ها را بردارید:\n' + '\n'.join(invalid), parent=self)
            return
        if not jobs:
            messagebox.showinfo('انتخاب شیت', 'حداقل یک شیت را انتخاب کنید.', parent=self)
            return
        if len(jobs) == 1 and save_visio:
            path, names = jobs[0]
            output = filedialog.asksaveasfilename(parent=self, title='ذخیره فایل Visio',
                initialfile=Path(path).stem + '.vsdx', defaultextension='.vsdx',
                filetypes=[('Visio drawing', '*.vsdx')], confirmoverwrite=False)
            if not output:
                return
            if Path(output).suffix.lower() != '.vsdx':
                messagebox.showerror('پسوند فایل', 'فایل خروجی باید پسوند .vsdx داشته باشد.', parent=self)
                return
            planned = [(path, names, output)]
        else:
            folder = filedialog.askdirectory(parent=self, title='انتخاب پوشه خروجی', mustexist=True)
            if not folder:
                return
            try:
                planned = plan_batch_outputs(jobs, folder, export_pdf=export_pdf)
            except OSError as exc:
                messagebox.showerror('خطای پوشه خروجی', str(exc), parent=self)
                return
        merge_output = None
        if self.merge_pdf.get():
            folder = Path(planned[0][2]).parent
            reserved = {str(Path(item[2]).with_suffix('.pdf').resolve()).casefold() for item in planned}
            candidate = folder / 'Merged.pdf'
            merge_output = filedialog.asksaveasfilename(parent=self,
                title='ذخیره PDF ادغام‌شده', initialdir=str(folder), initialfile=candidate.name,
                defaultextension='.pdf', filetypes=[('PDF', '*.pdf')], confirmoverwrite=False)
            if not merge_output:
                return
            if Path(merge_output).suffix.lower() != '.pdf':
                messagebox.showerror('پسوند فایل', 'پسوند فایل ادغام‌شده باید .pdf باشد.', parent=self)
                return
            if individual_pdf and str(Path(merge_output).resolve()).casefold() in reserved:
                messagebox.showerror('نام فایل', 'برای فایل ادغام‌شده نامی متفاوت از PDFهای تکی انتخاب کنید.', parent=self)
                return
        skipped = set()
        targets = []
        for _, _, output in planned:
            if save_visio:
                targets.append(Path(output))
            if individual_pdf:
                targets.append(Path(output).with_suffix('.pdf'))
        if merge_output:
            targets.append(Path(merge_output))
        for target in targets:
            if target.exists() and not self.confirm_replace(target):
                skipped.add(str(target.resolve()).casefold())
        if merge_output and str(Path(merge_output).resolve()).casefold() in skipped:
            merge_output = None
        if targets and all(str(p.resolve()).casefold() in skipped for p in targets):
            self.status.set('ذخیره همه خروجی‌ها لغو شد.')
            return
        self.busy = True
        self.set_enabled(False)
        direction = 'context'
        threading.Thread(target=self.worker, args=(planned, direction, save_visio, individual_pdf, merge_output, skipped), daemon=False).start()

    def set_enabled(self, enabled):
        controls = [self.browse, self.all_button, self.none_button, self.create_button, self.pdf_check, self.merge_check, self.visio_check, self.recheck_button, self.exporter_button]
        for w in controls + self.checks + self.group_buttons:
            w.configure(state='normal' if enabled else 'disabled')

    def worker(self, jobs, direction, save_visio, individual_pdf, merge_output, skipped):
        try:
            result = process_outputs(jobs, lambda msg: self.events.put(('status', msg)),
                                     direction, save_visio, individual_pdf, merge_output, skipped)
            self.events.put(('done', result))
        except Exception as exc:
            self.events.put(('error', str(exc)))

    def show_results(self, successes, failures, merge_note):
        result = tk.Toplevel(self)
        result.title('نتیجه ساخت فایل‌ها')
        result.geometry('720x420')
        ttk.Label(result, text=f'موفق: {len(successes)}    ناموفق: {len(failures)}', padding=12).pack(fill='x')
        body = ttk.Frame(result, padding=12)
        body.pack(fill='both', expand=True)
        scroll = ttk.Scrollbar(body)
        scroll.pack(side='right', fill='y')
        area = tk.Text(body, wrap='word', yscrollcommand=scroll.set)
        area.pack(fill='both', expand=True)
        scroll.configure(command=area.yview)
        if successes:
            area.insert('end', 'فایل‌های ساخته‌شده:\n' + '\n'.join(successes) + '\n\n')
        if merge_note:
            area.insert('end', merge_note + '\n\n')
        if failures:
            area.insert('end', 'فایل‌های ناموفق:\n')
            for path, error in failures:
                area.insert('end', f'{path}\n{error}\n\n')
        area.configure(state='disabled')
        ttk.Button(result, text='بستن', command=result.destroy).pack(pady=8)

    def poll(self):
        try:
            while True:
                event, payload = self.events.get_nowait()
                if event == 'status':
                    self.status.set(payload)
                else:
                    self.busy = False
                    self.set_enabled(True)
                    if event == 'done':
                        successes, failures, merge_note = payload
                        self.status.set(f'پایان پردازش — موفق: {len(successes)}، ناموفق: {len(failures)}')
                        self.show_results(successes, failures, merge_note)
                    else:
                        self.status.set('ساخت فایل انجام نشد؛ اطلاعات خطا را بررسی کنید.')
                        messagebox.showerror('خطا', payload, parent=self)
        except queue.Empty:
            pass
        self.after(150, self.poll)

    def close(self):
        if self.busy:
            messagebox.showinfo('در حال پردازش', 'لطفاً تا پایان ساخت فایل‌ها صبر کنید.', parent=self)
        else:
            self.destroy()

def main():
    if sys.platform != 'win32':
        print('This script requires Windows and installed desktop Visio.')
        return 1
    try:
        import openpyxl
        import pythoncom
        import win32com.client
    except ImportError:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror('Missing dependencies', 'Run in Terminal:\npython -m pip install openpyxl pywin32')
        root.destroy()
        return 1
    ImportWindow().mainloop()
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
